"""Genera i documenti SINTETICI del collaudo manuale (docs/collaudo/checklist.md).

Si lancia con scripts/genera-dati-collaudo.sh (container usa e getta con reportlab, python-docx, Pillow).
Tutti i contenuti sono inventati: comuni (Borgoverde, Castelrovo, Pievalta), persone, indirizzi e
protocolli non esistono. Gli indirizzi email usano domini di esempio (example.com / example.org).
Il finto motore sceglie lo scenario dal NOME del file: i nomi qui sotto sono quindi significativi.
"""
import io
import random
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

QUI = Path(__file__).resolve().parent
random.seed(42)
PROTOCOLLO = "0007101/2026"

ESPOSTO_CANI = [
    "Spett.le ATS Insubria - Dipartimento Veterinario",
    f"Prot. n. {PROTOCOLLO} del 03/03/2026",
    "",
    "Oggetto: esposto per cani tenuti in un recinto senza riparo",
    "",
    "Io sottoscritta Anna Esempio, residente a Borgoverde in via delle Prove 1,",
    "segnalo che nel cortile di via dei Collaudi 12 a Borgoverde sono tenuti",
    "cinque cani in un recinto senza riparo dalla pioggia. Uno dei cani zoppica",
    "da alcuni giorni. La situazione dura da circa due settimane.",
    "",
    "Chiedo un sopralluogo. Recapito: anna.esempio@example.com",
    "",
    "Borgoverde, 2 marzo 2026",
    "Anna Esempio",
]
RISCONTRO = [
    "ATS Insubria - Dipartimento Veterinario",
    f"Rif. Prot. {PROTOCOLLO}",
    "",
    "Gentile signora Anna Esempio,",
    "in riferimento al suo esposto sui cani tenuti in via dei Collaudi 12 a Borgoverde,",
    "le comunichiamo che il sopralluogo congiunto con la Polizia Locale si e' svolto",
    "il 20 marzo 2026. Il proprietario ha predisposto un riparo e il cane e' stato visitato.",
    "",
    "Cordiali saluti",
    "Il Dirigente del Servizio (firma di prova)",
]
EMAIL_PAGINE = [[
    "Da: Polizia Locale di Castelrovo <polizia.locale@example.org>",
    "Inviato: lunedi 9 marzo 2026 09:42",
    "A: protocollo@example.org",
    "Oggetto: I: segnalazione galline e oche in cortile",
    "",
    "Si inoltra per competenza la segnalazione ricevuta.",
    "Una pattuglia e' passata sul posto senza riuscire ad accedere.",
    "",
    "Il Comandante (firma di prova)",
    "",
    "-----Messaggio originale-----",
    "Da: Mario Prova <mario.prova@example.com>",
    "Inviato: domenica 8 marzo 2026 21:15",
    "A: Polizia Locale di Castelrovo",
    "Oggetto: segnalazione galline e oche in cortile",
], [
    "Buonasera,",
    "abito in via dei Gelsi di Prova 3 a Castelrovo. Nel cortile confinante",
    "sono tenute circa 40 galline e 6 oche a meno di 10 metri dalla mia finestra.",
    "Ci sono escrementi accumulati lungo la rete e l'odore e' molto forte,",
    "soprattutto la sera. Ho gia' avvisato i vigili due volte.",
    "",
    "Grazie, Mario Prova",
    "",
    "Questo messaggio e' riservato (clausola di esempio).",
]]


def pdf_testo(nome: str, righe: list[str]) -> None:
    c = canvas.Canvas(str(QUI / nome), pagesize=A4)
    c.setTitle("Documento sintetico di collaudo")
    y = A4[1] - 72
    for r in righe:
        c.setFont("Helvetica", 11)
        c.drawString(64, y, r)
        y -= 18
    c.showPage()
    c.save()


def immagine_pagina(righe: list[str], storta: float) -> Image.Image:
    """Pagina 'stampata e scansionata': solo immagine, leggermente storta e con rumore."""
    w, h = 1240, 1754   # A4 a 150 dpi
    img = Image.new("L", (w, h), 250)
    d = ImageDraw.Draw(img)
    font = ImageFont.load_default(size=26)   # font incorporato in Pillow, non serve nessun font di sistema
    y = 140
    for r in righe:
        d.text((110, y), r, fill=25, font=font)
        y += 42
    img = img.rotate(storta, fillcolor=250, resample=Image.BICUBIC)
    rumore = Image.effect_noise((w, h), 18).point(lambda v: 255 if v > 140 else 230)
    img = Image.blend(img, rumore.convert("L"), 0.12).filter(ImageFilter.GaussianBlur(0.6))
    return img.convert("RGB")


def main() -> None:
    # formati
    pdf_testo("C01_esposto_digitale.pdf", ESPOSTO_CANI)
    pagine = [immagine_pagina(p, s) for p, s in zip(EMAIL_PAGINE, (0.8, -0.6))]
    pagine[0].save(QUI / "C02_scansione_email_inoltrata.pdf", save_all=True, append_images=pagine[1:], resolution=150)
    doc = Document()
    doc.add_heading("Esposto per colonia felina in cortile", level=1)
    for r in ["Spett.le ATS Insubria,",
              "sono Lucia Collaudo, abito a Pievalta in viale degli Esempi 7.",
              "Nel cortile condominiale vive una colonia di circa quindici gatti; alcuni appaiono",
              "denutriti e uno ha una ferita a una zampa. Chiedo un controllo.",
              "Pievalta, 5 marzo 2026 - Lucia Collaudo (lucia.collaudo@example.com)"]:
        doc.add_paragraph(r)
    doc.save(QUI / "C03_esposto_gatti.docx")
    rtf = ("{\\rtf1\\ansi\\deff0{\\fonttbl{\\f0 Helvetica;}}\\f0\\fs22 "
           "Spett.le ATS Insubria\\par Oggetto: rumori notturni e odori da laboratorio\\par\\par "
           "Sono Paolo Esempio, abito a Castelrovo in via del Collaudo 21. Da un mese un laboratorio "
           "vicino lavora di notte con rumori persistenti fino alle 3 e odori molesti.\\par\\par "
           "Castelrovo, 6 marzo 2026 - Paolo Esempio\\par}")
    (QUI / "C04_esposto_rumori.rtf").write_text(rtf, encoding="ascii")
    (QUI / "C05_vuoto.pdf").write_bytes(b"")
    (QUI / "C06_danneggiato.pdf").write_bytes(
        b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n"
        + bytes(random.randrange(256) for _ in range(2048)))
    buf = io.BytesIO()
    Image.new("RGB", (320, 200), (200, 210, 220)).save(buf, "JPEG")
    (QUI / "C07_foto.jpg").write_bytes(buf.getvalue())
    (QUI / "C08_lettera.pages").write_bytes(b"documento sintetico in formato non ammesso\n")
    pdf_testo("C09_riscontro_ats.pdf", RISCONTRO)

    # un file per ogni scenario del finto motore (lo sceglie il nome del file)
    scenari = {
        "F03_foglio_manoscritto.pdf": "lettura incerta (foglio scritto a mano)",
        "F04_errore_formato.pdf": "formato non supportato (non si ritenta)",
        "F05_illeggibile.pdf": "documento illeggibile (non si ritenta)",
        "F06_scaduto.pdf": "tempo scaduto (tre tentativi, poi analisi non disponibile)",
        "F07_instabile.pdf": "errore al primo tentativo, riesce al secondo",
        "F08_occupato.pdf": "motore occupato una volta (429), poi riesce",
        "F09_invalido.pdf": "risultato fuori contratto (scartato)",
    }
    for nome, scenario in scenari.items():
        pdf_testo(nome, ["Documento SINTETICO di collaudo", "", f"Scenario del finto motore: {scenario}.",
                         "", "Esposto di prova: cane libero incustodito in via delle Prove, Borgoverde."])
    print("\n".join(sorted(p.name for p in QUI.iterdir() if p.suffix != ".py")))


if __name__ == "__main__":
    main()
